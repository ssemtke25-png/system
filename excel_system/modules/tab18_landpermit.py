"""
탭18(메뉴 ⑰): 토지거래허가 분기보고 취합
시군에서 받은 '토지거래계약 허가 현황' 파일들을 빈 서식에 붙여 취합본 1개로 만든다.

· 결과 파일 구성 : 맨 앞 '취합' 시트 + 올린 시군만 직제순(포항시 남구 → … → 울릉군)으로 시트 1장씩
· '취합' 합계 수식은 실제로 붙은 시군 시트만 더하도록 매번 새로 만든다
· 빈 서식(숫자 없음, 취합 / 서식 2장)은 이 파일 안에 들어 있어 별도 파일이 필요 없다
· 시군 파일은 스타일을 무시하고 셀 값만 직접 읽는다
  (다른 프로그램으로 저장돼 엑셀 라이브러리로 안 열리는 파일 대비)
· 검증 : 지역 인식 실패, 같은 지역 중복, 숫자 아닌 값(예: '해당없음') 0 처리,
         시군이 직접 입력한 합계와 세부값 불일치, 본기 > 누계, 분기 표시 불일치
"""
import io
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
import streamlit as st
from openpyxl.utils import column_index_from_string, get_column_letter

# 빈 서식(취합 / 서식 2장, 숫자 없음)을 파일 안에 넣어 둠 → 별도 서식 파일 없이 이 파일 하나로 동작
_TEMPLATE_B64 = """
UEsDBBQAAAAIANdMSF1Gx01IlQAAAM0AAAAQAAAAZG9jUHJvcHMvYXBwLnhtbE3PTQvCMAwG4L9SdreZih6kDkQ9ip68zy51hbYp
bYT67+0EP255ecgboi6JIia2mEXxLuRtMzLHDUDWI/o+y8qhiqHke64x3YGMsRoPpB8eA8OibdeAhTEMOMzit7Dp1C5GZ3XPlkJ3
sjpRJsPiWDQ6sScfq9wcChDneiU+ixNLOZcrBf+LU8sVU57mym/8ZAW/B7oXUEsDBBQAAAAIANdMSF0mKvq4NAEAAIwCAAARAAAA
ZG9jUHJvcHMvY29yZS54bWzFkk1OwzAQha+Csk/HdqA/VpoFIFZUqkQlEDvLmbYWcWLZrtLuewEugLgDx6KHwEnblAJ7ln7z/M2b
0aTScFlZnNrKoPUK3cVaF6Xj0oyjpfeGAzi5RC1cLzjKUJxXVgsfnnYBRsgXsUBghPRBoxe58AIaYGw6YpSlueTSovCVPeBz2eHN
yhYtLJeABWosvQPaoxBlu7ft58frbvuewonQ0Dxa7fYC5h2yVf/kthWIDs61U52rrutenbS+MASFp8n9QztvrErnRSkx/HKK+43B
cXTs/Jjc3M7uoowRymIyiuloRkacMM76z03Ws3ynwLrK1Vz9c2LWjymJyXBGCE+G/HLwLfExYJaGuyiE85ODcL3JVg5tCr/1o3Vq
VdlMG3YyaHbCQgfKkytOmp38MLXC+eFlX1BLAwQUAAAACADXTEhdl+vlUIAGAAAeIgAAEwAAAHhsL3RoZW1lL3RoZW1lMS54bWzt
Wltv2zYUfh+w/0DovZUvUuoEdYrYsZu1SRMkboc+0jItMaZEgaST+m1oMWDAhgHDumEvA/bWh2FbgRbYS/drsnXYuqF/YUeSL6JN
pU6b3dA4gC1S37nw3HhE5eq1eyFDR0RIyqO6Vb5cshCJPN6jkV+3bnfal2oWkgpHPcx4ROrWiEjr2vq771zFayogIUFAH8k1XLcC
peI125YeTGN5mcckgnt9LkKsYCh8uyfwMfANmV0plVbsENPIQhEOge1uv089gv788JMXjz6y1ifcWwy+IiWTCY+JAy8VmSdJsb1B
OfmRI9lkAh1hVrdAUI8fd8g9ZSGGpYIbdauUfix7/ao9JWKqgDZH104/Y7oxQW9QSemE350SltvO6pXNKf9Kxn8R12q1mq3ylF8K
wJ4HKy0vYJ12rdyY8MyBsstF3s2SW3J0fI5/dQG/2mg03FUNX53hnQV8rbTibFQ0vDPDu4v6NzaazRUN787wKwv49pXVFUfHp6CA
0WiwgE78OfXMFNLnbMsIrwG8NgmAGcrORVdGH6miWAvxIRdtAKTOxYpGSI1i0sce4Jo47AqKLRTjiEuYKFVK7VIVvpM/J71yEvF4
jeAcXTblyYWpRBMkPUFjVbduAFcrB3n57LuXz56gl88en9x/enL/x5MHD07u/2Ag3MKRnyd88eizP77+AP3+5JsXD78w42Ue/8v3
H//80+dmoMoDn3/5+Nenj59/9elv3z40wDcE7ubhHRoSiW6RY7TPQ1ibQQDpirNRdAJMNQocANIAbKlAA94aYWbCNYhuvDsCyoMJ
eH14qOl6EIihogbgzSDUgDucswYXxuXcTGTllzOMfLNwMczj9jE+Msluzrm2NYwhzqmJZTMgmpp7DLyNfRIRhZJ7fECIgewupZpd
d6gnuOR9he5S1MDUaJIO7Soz0RYNwS8jk4Lgas02O3dQgzMT+01ypCMhITAzsSRMM+N1PFQ4NGqMQ5ZHbmMVmJQ8GAlPM7hU4Gmf
MI5aPSKliWZXjDR1b2KoU0a377BRqCOFogMTchtznkdu8kEzwGFs1JlGQR77nhxAiGK0x5VRCa5nSDIGP+Co0N13KFFnS+vb1A/M
AZLcGQpTShCu5+OI9TExMd8QoVZYN6CGm6KjMfS10N4mhOFj3CME3X7PhOcxNyt9I4CqskVMtrmB9VhNxhGRBKXNjMGxVGohe0B8
XqDPzmiu8IxwFGJRxPnWQA+ZFuxtxlK6y7yBVkqpSJLWrMSuDPFSXPcCrIVVMpbmeB2J6Kw5BjSHr0FDzkwDhX1p23QwI+aA6WCK
tk3lFkiGZpIknVKyoZGuryftzA32XJMT0ui0jodRoJzreNyLjsfc8RRVlvk+pwj3P+xuNvEw2iOwoVw0NxfNzdvY3BTl8kVLc9HS
XLQ0/1hLM+ti7PwJT8olLDzu6VPGDtSIkW2Z9j8Scr/Xhsl0kBJNT5fiAC7H4jScL3B6jQRX71MVHAQ4BjHlVIIvx6x9iWIuoYOy
CnknN6D/UtmcOznNBDRWO7yXTVfzp5xTNunIl3lB1YTBssKqV95MWDkDLimt7JqluadKs3PWhE0F4eQQu7xSyURDnEAs9hK7Zwwm
bjl3F8kACujYR2XjQsrVJc1We7XVctJWq28mbRkn5cU5BeLcc/BSacFL9mI6skgfoWPQyq24FvJwXLf68LgCl2EM/GRSHDDzo7rl
qfFSXpnM8ws2h2W5VLhgTUQspNrEMsio0luTlwDRTP+K6yR2OJ8FGKrRclpUa+V/UQt73rWk3yeeKpiZDcf3+FARcRD0jlGXDcU+
Br2dLLp6VMJOUZkMBGSoMw48PfPHWTD/smGcHZjFAR7XpFrO9xk8vZ7qkI5y6tkFur/mUqrnuBT37V1KErnw9FftpecS0AYIjJIY
rVtcqIBDFYoD6rUFNA6pLNALQVokKiGWvDtNdCVHs7qV8ciKnB+ofeojQaHSqUAQsqfG63wFs3Ilv79OGI3rzFRdGWe/XXJEWCfJ
3pVk/RYKJtVkbIgUN+8025RdXb/9H+58nILO5/T2YCbIOUsv4uSKfm4rWH0zFc641VbMK664S2+1MTzDo+QLCjcVHpv1tx2+D95H
044SQSBeqo3TbzrZBZ1rucUlrP7eNmrmglqBv8+z+cwZu1pg7NPFvb6xXYOt3dNNbS+mqJ17kElHC/9CwbuHIHsTHo+GTMnsUPYe
PO81Jy+/gY89I13/C1BLAwQUAAAACADXTEhdknC8lqIJAABdMwAAGAAAAHhsL3dvcmtzaGVldHMvc2hlZXQxLnhtbL1ba1PjuBL9
fn+FK1u1M3PDJbaslx2ganiE8AiEJATCtwAGUpsHNzHDbm3tf7+yJIfYUctJoO5U7cbWabW6W31akm123ifTP2YvURQ7f46G49lu
6SWOX8NKZfbwEo36s+3JazQWyNNkOurH4nb6XJm9TqP+o+w0GlaQ69LKqD8Yl/Z2ZFtzurcT9+8PJsPJ1Jk+3++WauKfK/6VKns7
k7d4OBhHzakzexuN+tO/9qPh5H235JXShtbg+SVOGoT0a/85akfx9WtzKu4qc/2Pg1E0ng0mY2caPe2WfnrhnRckHaREdxC9zxau
ncTJ+8nkj+Tm5HG35Ca2RsPoIU5U9MXPr+ggGg6Fpp9MGPJfpVVcz0dNui5ep/prMjDCnfv+LBI+3wwe45fdEi85j9FT/20Ytybv
9Ui75G+TROPDZDiT/3felTTZZqTkPLzN4slIK0jiEf81jHZLLCg5o8FYNo36f+rIZHqjFXoj3dvP9eYrjY11b7zR2ET3phuNzXRv
ttHYXPcOcr2DbX+loLtp1N2Nhvfms4Zy/em2t5ICP1WQn7kVDUinziN5A1boTNPO+albcfB07lA+cbiggnMfzeLaQJK9OH/TJEJL
fmzzVWYSpb6gvC/BajOB5s6ouqBoLKvAYT/u7+1MJ+/OVHZN2I7I9lzpvAIIrVKjGD8R/akaqCCnQAfjpDS246mAB0J1vPf7b8Sn
FFfFr0cxF78Yezi5x0wUVHmPUNJOXEJY1REXmCEl6PpIN3ApgTkPqt9VFy+R4MTzq87fSWfKUNJCXeRW/xEXxOdu1UmwRIpKfSRA
Ess0/NipxCIWibmVB/GfiME8EEgHwl/IFTASSDaoQCRrSaJNIvsK8ZaRAxA5BJEjEKmByDGI1EHkBEROQeQMRM5BpAEiFyByCSJN
ELkCkRaItEGkoxDqA3mfZCnhMsGJG4jEd0JH5T6vbsmcJgGVV56PPGbNQn+dLPSVXRjkI8Z+QKspC7LDqnRVKgiDVQRcERL5Aav+
6/ffAt9D1V3561fL8jcwqT5QqrlrSHgNmTJeQ5vXmIKSYrC0psf0DWSCoToMncDQKQydwdA5DDVg6AKGLmGoCUNXMNSCoTYMdWDo
Goa6GsLL0E1hPhPEPLl6BESmD5S+t3D69uD0vVMQg9MXc+aq9EQkT8lMJcDrVAKsLCKGBUlD1LAiaYgZGKohbmAoLi4ZRDNTrfay
UqiKgXXFIEYeYpiHGJzyurYHqs75iqDtMIx/ogcJDLRVUGDIhjMNeZYpJ2i5MhkMOIcNaMDQBQxdwlAThq5gqAVDbRjqwNA1DHXh
mN/ASX0LJ3UPTuo7PVZ2R5fhI9F8ZIVkJDAZFUQt3JG7hoRDVCxapmVVqcDEku5ELvrzBRJSdaitcWFrGJOMIR5yPZOKo887VPs6
h44/71CdFFWTQodOvs6hU62Kbu7QGdmgXBNdrqku10z/Go08J/DmhIBl+6Iw0MRzqdpTc3FYlK5yMnfVUbvaD9+TsFLfr2rDDYZe
ErgaErDWXBVPZ35Hqs+cwh6ElaGeYQHQkTUY2oINbcOGdgoNFTPMF2cac8RUZD3qZQz1MKHzxISn/ho2tAsbeqMN/UTRuNUqgs9z
rKcT0bJ+F3HsjhSuHlSvHsEquzkKrm77CgoMXDvQkIFrhxoyrEhHsMIarPAYVlin4CJ7Ais81ZDB5TMFMb55tpwrFeQLKnJDqaKf
yJYLbY3t2F7g0OXXOdQsVFXo0FWxiiKHWl/nUPvzDnU+79D11znU1ao+sQm4gVl+C5OyB7P8TkMMrnhMVTyfb+PiisfUszcLx7nP
Pp7jmp5kKRUBF52e9o5Y+YbtVJ72dn4JmV+LZVLLBVKuxsq3ZrlDbZLrSsFjVu6ZBY9SQU8Ktq8b3+ts64z9MErXMsMnwids6xwQ
Pk5Vo7n0KdtqANL1VNr0RMqCnVqwsxTD0oILVr5i5Y45EOepLJGyl6zcYuVrs2wjlaVStsnKbVbummUvLPZdWrCmBbuyYC0L1rZg
HQt2bcG6KWY4st6kmGE5vbXo7KWYYWd2pzHPhcnL1yEvLyQvIeqNjdx3msjLM+Tl5RtuJi/PkJeXb81yh9okT3HymJd7ZsEjvkxe
vnXGAfLyJfLyrXNA+Jgvk5dvNQDpOreQ14KdWrAzniUvL1/xcscciHOeJS8vt3j52izb4Fny8nKbl7tm2QuLfZcWrGnBrixYy4K1
LVjHgl1bsC63kJdbyGvR2eMW8vJi8iaYPGyg7eKHVYnwTL53MpE3+c4mnL32H6Ld0us0mkXTX1Fpz/m3I4/nPJD7FiyP5+KY7tP5
YdLyYNvz1jHP28A85z9OfouWFCFhmHwf5CUv4RYAjtTxHXPsqhdGVL0zJkg9qmdeIHtm34Ulb6BFDKpLY33/2NylYxCXEfkyCrMA
4Wp6FLe9CfTQOmFCm4bJvA2dh8v7cI64VMUHDAOoi4nQqjglu7dEwBPKv0udVEm4SmJJRfoWzxJ79YGAyD9qVgFboZ5yiEx28eI8
Z1PbOkn+OpPkbzJJjiPDJCz74XzLP2xS2Ypck98f85VmJwtELL+lT9Yc+QkFlh9XiODiVDTzwcacNPlRVLC2jCZhSt0M5eYphCgK
Fi3BjFDzwKlF4MD55PnwRrincybj+ILfxTFLUsDN5AZL3yab7bElCV4nSfBmSeLMMyVLqDWCUsTFVYKzFvW2FuO5qaH/BzsLS4Rt
8sk6k082K+Pf8vVRkTytytmnzllJ2+qVDc8PA9mXiVvw+cWiYSuvIx+VjHCkZjwlIdKPBRjRSSNaqOpCA12eMm/8vy2WpZwBRDRo
tdmPBD6SdMVQ2RKCrpMQdOOEyE5E4jYNWJrLKsxrTeUq251cFPILvi0qbJ2osM1pYnwXb4qOWXIxCplUM2eplwrohPcV4KqxMFf3
HxvMFQO7tJWyRZavE1l1vEDAY0vrfjvA8pC/sE30iZ+r235mU5HtsrAbEF3SF1PYfpYI1nEuWEib7MmpsvBh7CiaPssP7GfOw+Rt
HMtvcBea1Vf8RyQ8kp/m5tpvSHhjar8iYZsY2nsk7JnkL0jYNMnXSFgzycu/KjC03/nhnVHeD3+a2vf98BAb2m9JeGuSPyHhian9
mITHpvYzHHZN+vdJuG+SPyThoTHOftgz6Tnyw65vaK+TsG7S0yFh1xTnUxKeGu0nYcMkX8fhqdEeHB6b2jsovEOG9gMSHshxKx95
qP60pNGfPg/GM2cYPYmcdJOEn6pMl9fx5FVeJV+qT2LBg/TuJeo/RtPkTmT+02QSpzfJIPO/qNn7H1BLAwQUAAAACADXTEhd4Tqv
KyQKAAApNAAAGAAAAHhsL3dvcmtzaGVldHMvc2hlZXQyLnhtbL1bW3PiuBJ+P7/CxVbtzBzYYMuSJZskVZMLkARyAUJC3pjEmVAL
mGOcyc5s7X9fWZIBG7UNJHUeEkBfq9Xd6q8lW/b+WxD+OX/x/cj4azKezg9KL1E086rV+eOLPxnO94KZP+XIcxBOhhH/GX6vzmeh
P3wSnSbjKjJNpzoZjqalw33Rdh0e7gev0Xg09a9DY/46mQzDn0f+OHg7KFmlpKEz+v4SxQ3Vw/3Z8Lvf9aPbGZd/HkW94Jo3KKy6
0Pk0mvjT+SiYGqH/fFD6ankPlhuLCIn+yH+br3w35i/BWyMcPbW4Hdwts2REw29df+w/Rv6TsORXEEy6j8Oxfxk7N+ZtprnS2o1V
tYY/uS+xwoOSy9E4XN+C4M+45ewpVsvHFEpjw4b844d/7I+5sq8nmDv7P2kr/77wJe66+j2xui5CzEPwbTj3j4Px3egpejkosZLx
5D8PX8dRJ3hr+ips9h6JNT4G47n4b7xJabJHScl4fJ1HwUQpiGMe/RzzgFK3ZExGU9E0Gf6lIpzqjTbojVRvO9ObbTQ2Vr3xTmMT
1dvZaWyqetOdxmaqt5vp7e7ZGwXdTKJu7jS8tZg1lOnv7FkbKbATBdmZ29CAZOoskjVgg85O0jk7dRsOnswdyiYO41QwvvnzqD4S
BaU4f5MkQmt+7LFNZhIlvqCsL+5mM4EWzohMrEoaiypwMoyGh/th8GaEomvMdkT2FkoXFYBrFRr5+LHoV9ngcHJydDSNy283Cjk8
4qqjw99/I7bj4Br/tBzM+CfGFo5/Y2qapviNUNxOTEJozeBfMEVS0LSRamBCAjPm1j7/vdRFTMdxa/8Yxt/iO0VxN8dEZu0f/oXY
zKwJDDPeQ6gkLhJYquHLfjXi4Ygtrj7yPx6GRSyQioW9ki5gMJBokLGIF6ZYm0COJGKtI8cgcgIipyBSB5EGiDRB5AxEzkHkAkRa
INIGkUsQuQKRaxC5AZEOiHRBpCcRxwZS/3OcgkzkODFdnq+GZ8j0Z7WKSGviOuKbZSOL5mahvU0W2tIuDFISY9sV2S9YkB5WpqtU
QSiswmWSk8h2ae0/v//m2haqHYhPu1YWn65O9bFUzUxNwitIl/EK2r3MFFQVjaV1NaatIRMMNWHoDIbOYegChlow1IahSxi6gqFr
GLqBoQ4MdWGoB0O3MNRXEF6H7grzmSBqidXDJSJ9oPS9h9N3AKfvg4QonL6YUVOmJyJZSqYqAd6mEmBpEdEsSApyNCuSgqiGoQpi
Gobi4pJBFDPlgi8qhawYWFUMouUhhnmIwSlvKnug6pytCMoOzfhnahBXQ1sJuZpsuFCQlTPlBK1XJo0BLdiANgxdwtAVDF3D0A0M
dWCoC0M9GLqFoT4c8zs4qe/hpB7ASf2gxkrv6FJ8JIqPtJCMBCajhJwc7ohdQ8whhy9aumVVqsAkJ92JWPQXCySk6kRZY8LWUCoY
QyxkWjoVp+93qP5xDjXe71CTFFWTQofOPs6hc6XK2d2hC7JDuSaqXDuqXFP1qTWyReDNCQHL9mVhoIllOnJPzfj1onCVkYWrhtzV
Ln2Pw+rYdk0ZrjH0isDVkIC15qZ4OrM7UnXNye1BWBpqaRYAFVmNoR3Y0C5saK/QUD7DbHWmMUNURtZyrJShFibOIjHhqb+FDe3D
ht4pQ99RNO6VCvf9HBuoRMxZv4s49kAKVw9HrR7uJrs5B1zdjiTkarh2rCAN104UpFmRTmGFdVhhA1bYdMBF9gxWeK4gjcsXEqJs
92xpSRXkAypyW6py3pEtl8qavMv2AoeuPs6h60JVhQ7dFKsocqjzcQ513+9Q7/0O3X6cQ32l6h2bgDuY5fcwKQcwyx8UROGKR2XF
s9keLq54VN57y+E4s+nyPq7uTpZSYVq81/PhKS3f0f3q8+H+Dy70Y7VOSkHXFXJ1Wr7Xy50kCpEQbNDyQC94mh65e9v+3KSVC/pF
K11PDR8Ln9FKCxBupG2Ipc9ppQ1INxNp3S0phVmae9bnCaa7KZXoxMKCS1q+oeWePhCtRJYI2Sta7tDyrV62ncg6Qvaalru03NfL
Xib2aXL0KsE0SXqdg93kYJ0crJuD9XKw2xysn2Ca1e8uwTTXs/c5OgcJprvYVRiBycu2IS8rJC8h8sRG7Dt15GVp8rLyHdOTl2XZ
U2eVe6bnwwlbY0+DVQaA9ClbpzGrXADS9TVDzlilBQg31g05Z5U2IN1kOTRmOTRmOTRmaRqz8g0r9/QxbrE0jVm5w8q3etk2S9OY
lbus3NfLXrIcGrMcGudgNzlYJwfr5mC9HOw2B+uzHBqzHBrn6BywHBqzQhrHjwSIyw60V3zbKhaeixMoHY3jx3e8+Wz46B+UZqE/
98MffunQ+K8hLtSZK3YwWFyo8wt221lcVubc4rasbcyzdjDP+MPIbtbicsQNEydDVnwctwIwJC/kMcOmPDpy5OkxQfKmPbVc0TN9
KhafRfMY1NbG+rzc5iVjEJMScSyFqYtwLbkozzsTtNA2YUK7hkm/IV2Ey1o6R0xHxgcMA6iL8tDKOMX7uFiA80XEyXQcKWFKiTUV
yXleTuwRsmT+OXoVsBXyfgfPZBOvznM6tXMnyd5mkuxdJskwRJi4ZV+MT9nbTjJbkanzezlfSXZSl8fyU3KPzRAPU2CrJoOLE9HU
0xsL0mRHkcGqaE3CjmOmKLdIIeQgd9USTImjHzixCBw4mzxLb7h7KmdSjq/4XRyzOAXMVG7Q5FxZb09ekuBtkgTvliTGIlPShNoi
KEVc3CQ4W1GvshrPXQ39P9hZWCLyJp9sM/lktzL+KVsfJcmTqpy+/5yWzFu90uH5oiH7OnELHsRYNWzjdWRZyQhDcsYTEiJ1g4AS
lTS8xZFdHFeVp9TZ/6fVspQxgPAGpTb9uMAySTcMVV5CONskhLNzQqQnInbbcWmSyzLMW03lJtudTBSyC35eVOg2UaG700R7Kq+L
jl5yNQqpVNNnqZUIqIS3JWDKsTCTv5cbzA0Du7aVyoss2yay8uoCATcwc/fbLhaX+yvbRJvYmbptpzYV6S4ruwHeJTmiwvnXEu42
zrkraZO+cqquPCU78cPv4mn7ufEYvE4j8UDuSrN8UeCUeKfiOd1M+x3x7nTtN8TrEk37gHgDnfwl8a518nXi1XXy4sUFTfuD7T1o
5W3vq679yPbkawWZ9nvi3evkz4h3pmtvEK+ha7/AXl+n/4h4Rzr5E+KdaONsewOdnlPb69ua9ibxmjo9PeL1dXE+J9651n7itXXy
Teyda+3BXkPX3kPeA9K0HxPvWIxbXebh4f4sHE2jq1n8OsjceAnC0a9gGg3Hx/408kP54ol646U9DL+PuNDYf+aZa4qDx1ASQv6I
gln8NX4i/VsQcb4kv1784ZMfxr84Q56DIEp+LN+leZ0ZfGw+6DC25KA0Hk6f5o/DmV8yZvx/2B398uM3GYx5/KqLPEQQr94knBTv
KFQXbwcd/gtQSwMEFAAAAAgA10xIXXE0uhiKCAAAsIsAAA0AAAB4bC9zdHlsZXMueG1s7V1Lj9s2EP4rggL0UHSrp2WrWS+62Y2R
Ak0RdPfQoi0CrS3bQvRwZTnx5tQHULS9FeilQI9BgAI5Fcktv6bHZvc/lKRkS7ZJWQ9TJrfrRWKJ5Mx8nBkORzJFHU6jS9c+G9t2
JMw91592xXEUTT6SpGl/bHvW9MNgYvugZhiEnhWB03AkTSehbQ2mkMhzJVWWDcmzHF88OvRnXs+LpkI/mPlRV1SXRUL89cmgKyqG
Lgoxu5NgYHfFxwfvC3c+uHNHfnxw9+vMMSx/79tZEN09iL9Q2cePD0QJy7a1zjam+ueHv+KDrBRiFYaKiEFKunt0OAz8tNeKIsYl
AKXl2cJTy+2K/755/e7l26u/X0PC/tgKp0DjqEZRTVg2tDzHvYyLNFgwfZ7UK0hWzJB5tp0crg++vP7+1buXr65+fXH9488Pi/BG
IC8yxws5qkxHKxodtsZO2eIUoqh0kO9Wz1nk/cANQiEcXXTFHvjI4EOxQ9uk7XY4rEiT5Vxp6AsGEcd1l0FEF+OCo8OJFUV26PfA
CaJBhRtVQnJ8fjkBsW8UWpeK2hILE0wD1xlAkaMThNzxB/bcBmFVj0fgejGI4ZB3hl99SVpZSegLKO4iCAd2uFRdC8bfuOzo0LWH
EaAPndEYfkfBBAoJoijwwMHAsUaBbyHFLihwlAKaJsHcaDmhuLDuOkTUEoko1By0K4GkEEvYkFvEEitAKqouGqMUaCviQu3jpsX7
lsuMa7co1MNGEFPW8V79h7I1uOibtH/INecbJjDz5BpFhzSDrkFZzbme1xDmojGcR8zs5C170XIzySEVyIylAXsxx/+hbzwEWMbS
fx7txsQkQ2kcsTPJ7BFxvdyNBXNzMd0xoOZ6kDnU8g6TTYmdK51qA1xi5x4Ui1AY8NoqaqkogUr0rjY0Soy9ktNDGX1KDeqpajdY
AMNEmsDAVLquPgqRr66IBlRPfX7lIuegGJaagsLEhSlF/dGGwsC9v91hSQ6mgNZ23TPI7YvhcvWAAXjOh5mFZTJcVubHa8zAoeO6
SWnMJjmxJhP38th1Rr5no0VggI21OBWe2mHk9OHKjz44tUOEcT5ck6QrzYlSGxOV1Z/coAJpy1KpypKyvhl7asZJFcXMcdNSaISJ
8zSI7s3AePFRE7jM0X4U2kNnjs7nw5i2luV1MqSi8iUy9xZV7gZ1dY6D0HkOpEGFwsAnlldwOwWpZkEqVEAmoDZhCs9Ca3Juz6Nk
jV8lzCqHmDXWMCfLrDdQa1nUeg3USgXU03Ho+E/Og55TS9kt1pRdALPBGuZCDtJm0EEKAe9wBFxmfUgWAm7yClypM5E3GwVXce9u
bi+WbhTSJk8RQ2Y9YhRJ8JrOlkorXOE1RJOAM+/iJODMzy2kWxFMhuhCyLlVOYvTeSYcdpi9KC80EmtF7b3GEK6G4gpyjXnkVDK9
vUYR9nVOQs7TJM++s2Qit8nI3Wn6iYfOa+KxV+B1Eo9mgJN+X+HAWUjQ2fcWEnIW3YXH28E39LY7j78l8fibHXO/fxXRM4+DUKmT
mO5tFDbs0pxEOF4cl5PgWxkmJzh5sTqXYZVDzDzmNlzmCXV+qdob6DoXRnvzaDq37ymD5nEcshyjtUZ/26k8/+WsaaUAEy2Fr+6j
DV81VMbZcKitjLPh6MrLhdduAmrDSUJlnHXumjZ5uciLs/Jy+c2Jfyq8+CcLdk/2D6d3xaU2CbTWHaJmkdbJTHaFlL1bg6UmUtJv
Y/JOnn8icdcMquxhFrAD9iR76oxdTZDUoJcAVUHLOmUf0emyb9Fl36HL3qRrWrrKUWsMf37msCJ6VumasU2XPV0n1HYTxUnsW3TD
l14D/Q1fB8zTwvYbsoCZpyWp3D7pI/PlKzfuwV3m4wq3TxzfrtFnBjiTa/Rv3qMonD3+c/sMITNezryv3G4fwQpw5ichxjc8kJK9
zTLb8a1sxrcsFeALGbvi9W9/Xr34Lr3uFi5mjhs5/vIqfJ3g6pe3gEb4Sv4mRZ8lMnBE1z+9uf7j9yyRukLUxhIhaIK6INHy4aQN
9XwIacP4lY6poo4OB/PMtnDwgZYBVGr8gsTCL2FMXlRpKB31FKnjglSBfxEjkllDsKkft057GMFpBVbwYktH0na5m/C3bVq+SbHY
SjLdPzLT2wC+wDP7UtPFOz131KZiB0thXbRd1YpnD5yZV0svK5TpS1DXxJJ3x90KAWPMwrAX8QmUpeFLInSnklMf37/Xu3+Mceq0
gs5o2oXgDWdZ0Wf8yRhzf/HmFugugOLjTP/J+kDKwCwRaFjs6cYckNvZusGuWh9P26e9no7pY1qR20fiFtubHdz2iolNii2b5hMc
ZvvUi9H6IlYXJFs0z0b2gqQpwYat4f8g34qsC9dezVJBdjawh9bMjc6XlV0xPX6I5iRz2eoRTLGTVunxp1CfCsx4UxFJJvhZEHqg
C3GRIhwIx32YVAuLdB2lwQmY9gqD+66NMvQIudmzceDaCBfAkqTCKEclEYxtC2jg8+BZ2l7Nax8FQHMrzZW85kMnnEbAsWeen1LI
eRSutUFgbpUAAJ1FoTNJe90piGqNrB378ZIOYypkRaylUNMaljLKGapVyk56WTPlug3GSrluQ7BSUd9ZI5M3rJQ5mSZjH3wne8af
JKcgGGRirCynMXy9Jo0XmzUkmrgOXwPrSHJICEg0MRVJzk3qT4fYn7iOhK2DrekQaTpEmpgKV3OC/khy8DQm+OB7apqaZhgkjZ6c
YBGckPRmGPAfnhsJG6QgyYGSyumabG2yh+T7AcmmeR5C6inZE0k9Jesa1uD1BilME29tkhxIQbICyXegfLwc6FN4Gk2DViVhI41g
co1pkmqgL+J91DAI2jHgH94+pFGiaaaJr4F1eASaRqqBo5FcQ0IAMZBqNDSZSmvzkbSYpySUup6NbTs6+g9QSwMEFAAAAAgA10xI
XZeKuxzAAAAAEwIAAAsAAABfcmVscy8ucmVsc52SuW7DMAxAf8XQnjAH0CGIM2XxFgT5AVaiD9gSBYpFnb+v2qVxkAsZeT08Etwe
aUDtOKS2i6kY/RBSaVrVuAFItiWPac6RQq7ULB41h9JARNtjQ7BaLD5ALhlmt71kFqdzpFeIXNedpT3bL09Bb4CvOkxxQmlISzMO
8M3SfzL38ww1ReVKI5VbGnjT5f524EnRoSJYFppFydOiHaV/Hcf2kNPpr2MitHpb6PlxaFQKjtxjJYxxYrT+NYLJD+x+AFBLAwQU
AAAACADXTEhd5xFsG3UBAADhAgAADwAAAHhsL3dvcmtib29rLnhtbLWSy07DMBBFfyXyHpJGPNSq6YaKh4SgAkTXbjxpRvgR2dOm
7Zoda3awhy/gp8o/4LgKFCEhNqzse8eauT52vzb2bmLMXbRQUruMlURVL45dXoLibtdUoH2lMFZx8tJOY1dZ4MKVAKRknCbJQaw4
ajbot71GNh70m80tQu2+/EZGc3Q4QYm0zFjYS2CRQo0KVyAylrDIlaY+NRZXRhOX17k1UmassyncgiXMf9jXTZ4bPnHBWYxRC1Nn
bKeT+obL77IOaoyCyoyl3WTv0zsFnJbkW+wfNgeJT644ocnYQeJlgdZRGBRi8pxwDn5mGMlnZI5REtghJzixZlahnjYlDyPeohHI
tesGe8/+BbwpCsxhaPKZAk0b8hZkE1C7EivHIs0VZGz99vL++Npw8RPOxIYR+VRbxG0PfcGeiZDvH7PcP60fnreypL9kSQOrFpCA
AjWIC9/Hed+/eT6yUbOEO3U7Sdr1jzKT8sh7l/rccNHybr/c4ANQSwMEFAAAAAgA10xIXY33LFq0AAAAiQIAABoAAAB4bC9fcmVs
cy93b3JrYm9vay54bWwucmVsc8WSTQqDMBBGrxJygI7a0kVRV924LV4g6PiD0YTMlOrta3WhgS66ka7CNyHvezCJH6gVt2agprUk
xl4PlMiG2d4AqGiwV3QyFof5pjKuVzxHV4NVRadqhCgIruD2DJnGe6bIJ4u/EE1VtQXeTfHsceAvYHgZ11GDyFLkytXIiYRRb2OC
5QhPM1mKrEyky8pQwr+FIk8oOlCIeNJIm82avfrzgfU8v8WtfYnr0N/J5eMA3s9L31BLAwQUAAAACADXTEhdbqckvB4BAABXBAAA
EwAAAFtDb250ZW50X1R5cGVzXS54bWzFlM9OwzAMxl+lynVqMnbggNZdgCvswAuE1l2j5p9ib3Rvj9tuk0CjYioSl0aN7e/n+Iuy
fjtGwKxz1mMhGqL4oBSWDTiNMkTwHKlDcpr4N+1U1GWrd6BWy+W9KoMn8JRTryE26yeo9d5S9tzxNprgC5HAosgex8SeVQgdozWl
Jo6rg6++UfITQXLlkIONibjgBKGuEvrIz4BT3esBUjIVZFud6EU7zlKdVUhHCyinJa70GOralFCFcu+4RGJMoCtsAMhZOYoupsnE
E4bxezebP8hMATlzm0JEdizB7bizJX11HlkIEpnpI16ILD37fNC7XUH1SzaP9yOkdvAD1bDMn/FXjy/6N/ax+sc+3kNo//qq96t0
2vgzXw3vyeYTUEsBAhQDFAAAAAgA10xIXUbHTUiVAAAAzQAAABAAAAAAAAAAAAAAAIABAAAAAGRvY1Byb3BzL2FwcC54bWxQSwEC
FAMUAAAACADXTEhdJir6uDQBAACMAgAAEQAAAAAAAAAAAAAAgAHDAAAAZG9jUHJvcHMvY29yZS54bWxQSwECFAMUAAAACADXTEhd
l+vlUIAGAAAeIgAAEwAAAAAAAAAAAAAAgAEmAgAAeGwvdGhlbWUvdGhlbWUxLnhtbFBLAQIUAxQAAAAIANdMSF2ScLyWogkAAF0z
AAAYAAAAAAAAAAAAAACAgdcIAAB4bC93b3Jrc2hlZXRzL3NoZWV0MS54bWxQSwECFAMUAAAACADXTEhd4TqvKyQKAAApNAAAGAAA
AAAAAAAAAAAAgIGvEgAAeGwvd29ya3NoZWV0cy9zaGVldDIueG1sUEsBAhQDFAAAAAgA10xIXXE0uhiKCAAAsIsAAA0AAAAAAAAA
AAAAAIABCR0AAHhsL3N0eWxlcy54bWxQSwECFAMUAAAACADXTEhdl4q7HMAAAAATAgAACwAAAAAAAAAAAAAAgAG+JQAAX3JlbHMv
LnJlbHNQSwECFAMUAAAACADXTEhd5xFsG3UBAADhAgAADwAAAAAAAAAAAAAAgAGnJgAAeGwvd29ya2Jvb2sueG1sUEsBAhQDFAAA
AAgA10xIXY33LFq0AAAAiQIAABoAAAAAAAAAAAAAAIABSSgAAHhsL19yZWxzL3dvcmtib29rLnhtbC5yZWxzUEsBAhQDFAAAAAgA
10xIXW6nJLweAQAAVwQAABMAAAAAAAAAAAAAAIABNSkAAFtDb250ZW50X1R5cGVzXS54bWxQSwUGAAAAAAoACgCEAgAAhCoAAAAA
"""


def _template_bytes():
    import base64
    return base64.b64decode("".join(_TEMPLATE_B64.split()))

# 직제순 : (지역키, 시트 이름, 공식 명칭)
ORDER = [
    ("포항남", "포항남구", "포항시 남구"), ("포항북", "포항북구", "포항시 북구"),
    ("경주", "경주", "경주시"), ("김천", "김천", "김천시"), ("안동", "안동", "안동시"),
    ("구미", "구미", "구미시"), ("영주", "영주", "영주시"), ("영천", "영천", "영천시"),
    ("상주", "상주", "상주시"), ("문경", "문경", "문경시"), ("경산", "경산", "경산시"),
    ("의성", "의성", "의성군"), ("청송", "청송", "청송군"), ("영양", "영양", "영양군"),
    ("영덕", "영덕", "영덕군"), ("청도", "청도", "청도군"), ("고령", "고령", "고령군"),
    ("성주", "성주", "성주군"), ("칠곡", "칠곡", "칠곡군"), ("예천", "예천", "예천군"),
    ("봉화", "봉화", "봉화군"), ("울진", "울진", "울진군"), ("울릉", "울릉", "울릉군"),
]
KEY_INFO = {k: (i, sheet, full) for i, (k, sheet, full) in enumerate(ORDER)}

PERIOD_LABELS = ("본기", "당월", "당분기", "금분기", "이번분기")
CUM_LABELS = ("누계",)
NOTE_COL = "Z"   # 비고 칸 (글자 그대로 복사)

# 시군이 합계 칸에 숫자를 직접 적어 보낸 경우 검산할 관계식 (서식의 수식과 같음)
CHECKS = [
    ("K", ("N", "Q", "T"), "불허가 소계 건수"), ("L", ("O", "R", "U"), "불허가 소계 필지수"),
    ("M", ("P", "S", "V"), "불허가 소계 면적"),
    ("E", ("H", "K"), "허가현황 소계 건수"), ("F", ("I", "L"), "허가현황 소계 필지수"),
    ("G", ("J", "M"), "허가현황 소계 면적"),
    ("B", ("E", "W"), "신청 건수"), ("C", ("F", "X"), "신청 필지수"), ("D", ("G", "Y"), "신청 면적"),
]

_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


# ── 시군 파일 읽기 (스타일 무시, 첫 시트 셀 값만) ──────────────────────
def read_first_sheet(data: bytes):
    """반환: {셀주소: (값, 수식여부)}"""
    z = zipfile.ZipFile(io.BytesIO(data))
    names = z.namelist()
    shared = []
    if "xl/sharedStrings.xml" in names:
        root = ET.fromstring(z.read("xl/sharedStrings.xml"))
        for si in root.findall(f"{{{_NS}}}si"):
            shared.append("".join(t.text or "" for t in si.iter(f"{{{_NS}}}t")))
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    first = wb.find(f"{{{_NS}}}sheets/{{{_NS}}}sheet")
    rid = first.get(f"{{{_RNS}}}id")
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    target = next(r.get("Target") for r in rels if r.get("Id") == rid).lstrip("/")
    if not target.startswith("xl/"):
        target = "xl/" + target
    sheet = ET.fromstring(z.read(target))
    cells = {}
    for c in sheet.iter(f"{{{_NS}}}c"):
        ref, typ = c.get("r"), c.get("t")
        v, f = c.find(f"{{{_NS}}}v"), c.find(f"{{{_NS}}}f")
        if typ == "s" and v is not None:
            val = shared[int(v.text)]
        elif typ == "inlineStr":
            val = "".join(x.text or "" for x in c.iter(f"{{{_NS}}}t"))
        elif v is not None and v.text is not None:
            try:
                num = float(v.text)
                val = int(num) if num.is_integer() else num
            except ValueError:
                val = v.text
        else:
            val = None
        if val is not None:
            cells[ref] = (val, f is not None)
    return cells


def _to_number(v):
    """숫자로 바꿀 수 있으면 숫자, 아니면 None"""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(",", "")
    if s in ("", "-"):
        return 0
    try:
        n = float(s)
        return int(n) if n.is_integer() else n
    except ValueError:
        return None


def _find_rows(cells):
    period = cum = None
    for ref, (val, _) in cells.items():
        m = re.fullmatch(r"A(\d+)", ref)
        if not m or not isinstance(val, str):
            continue
        label = re.sub(r"\s+", "", val)
        if period is None and label in PERIOD_LABELS:
            period = int(m.group(1))
        elif cum is None and label in CUM_LABELS:
            cum = int(m.group(1))
    return period, cum


def _quarter(text):
    m = re.search(r"(\d{4})\s*년\s*(\d)\s*(?:/\s*4)?\s*분기", str(text or ""))
    return (int(m.group(1)), int(m.group(2))) if m else None


def _find_key(text):
    """글자 속에서 시군을 찾는다. 포항은 남구/북구까지 있어야 인식.
    반환: 지역키 / '포항'(구 구분 불가) / None"""
    t = re.sub(r"\s+", "", str(text or ""))
    if "포항" in t:
        if "남구" in t:
            return "포항남"
        if "북구" in t:
            return "포항북"
        return "포항"
    for key, _, _ in ORDER:
        if not key.startswith("포항") and key in t:
            return key
    return None


def _region(fname, title):
    """파일명 → 제목(A1) 순서로 시군을 찾는다.
    (파일명이 '2. 토지거래허가 분기보고_경주'처럼 번호 뒤에 시군명이 바로 오지 않아도 인식)"""
    k1 = _find_key(Path(fname).stem)
    if k1 in KEY_INFO:
        return k1
    k2 = _find_key(title)
    if k2 in KEY_INFO:
        return k2
    return None


# ── 취합본 만들기 ─────────────────────────────────────────────────────
def build(files, year, quarter):
    """files: [(파일명, bytes)] → (결과 bytes 또는 None, 처리목록, 경고목록, 오류목록)
    오류(합계 불일치)가 하나라도 있으면 취합본을 만들지 않고 결과 bytes 는 None."""
    wb = openpyxl.load_workbook(io.BytesIO(_template_bytes()))
    form = wb["서식"]
    total = wb["취합"]
    isf = lambda v: isinstance(v, str) and v.startswith("=")

    # 서식의 입력칸(수식이 아닌 칸) 목록 : (행, 열문자)
    form_period, form_cum = 7, 8
    input_cols = [get_column_letter(c) for c in range(2, 27)
                  if not isf(form.cell(form_period, c).value)]

    log, warns, errors, used = [], [], [], {}
    parsed = []
    for fname, data in files:
        try:
            cells = read_first_sheet(data)
        except Exception as e:
            warns.append({"유형": "파일 읽기 실패", "파일": fname, "설명": f"엑셀 파일로 읽을 수 없음 ({e})"})
            continue
        title = cells.get("A1", ("", False))[0]
        key = _region(fname, title)
        if key is None:
            warns.append({"유형": "지역 인식 실패", "파일": fname,
                          "설명": "파일명·제목에서 시군(포항은 남구/북구까지)을 찾지 못해 제외"})
            continue
        if key in used:
            warns.append({"유형": "같은 지역 중복", "파일": fname,
                          "설명": f"{KEY_INFO[key][2]} 파일이 이미 있음({used[key]}) → 이 파일은 제외"})
            continue
        used[key] = fname
        parsed.append((KEY_INFO[key][0], key, fname, cells, title))

    parsed.sort(key=lambda x: x[0])   # 직제순
    sheet_names = []
    for _, key, fname, cells, title in parsed:
        _, sname, full = KEY_INFO[key]
        ws = wb.copy_worksheet(form)
        ws.title = sname
        sheet_names.append(sname)
        foreign = "외국인" in str(title)
        ws["A1"] = f"{'외국인 ' if foreign else ''}토지거래계약 허가 현황({full}  {year}년  {quarter}분기)"

        q = _quarter(title)
        if q and q != (year, quarter):
            warns.append({"유형": "분기 표시 다름", "파일": fname,
                          "설명": f"파일 제목은 {q[0]}년 {q[1]}분기 (선택한 분기: {year}년 {quarter}분기) — 지난 파일인지 확인"})

        src_period, src_cum = _find_rows(cells)
        if src_period is None or src_cum is None:
            warns.append({"유형": "양식 다름", "파일": fname,
                          "설명": "'본기(당월)' 또는 '누계' 줄을 찾지 못해 숫자를 넣지 못함"})
            log.append({"순서": len(sheet_names), "시군": full, "시트": sname, "파일": fname, "채운 칸": 0})
            continue

        filled, texts, got = 0, set(), {}
        for src_row, dst_row in ((src_period, form_period), (src_cum, form_cum)):
            for col in input_cols:
                raw = cells.get(f"{col}{src_row}")
                if raw is None:
                    continue
                val, _ = raw
                if col == NOTE_COL:
                    ws[f"{col}{dst_row}"] = val
                    continue
                num = _to_number(val)
                if num is None:
                    texts.add(str(val).strip())
                    num = 0
                ws[f"{col}{dst_row}"] = num
                got[(col, dst_row)] = num
                filled += 1
        log.append({"순서": len(sheet_names), "시군": full, "시트": sname, "파일": fname, "채운 칸": filled})

        if texts:
            # 한 칸에 한 글자씩 적은 경우(예: 해·당·없·음)도 원래 문구로 보여주기 위해 줄 전체 글자를 모은다
            row_txt = []
            for src_row in (src_period, src_cum):
                s = "".join(str(cells[f"{get_column_letter(c)}{src_row}"][0]).strip()
                            for c in range(2, 26)
                            if f"{get_column_letter(c)}{src_row}" in cells
                            and _to_number(cells[f"{get_column_letter(c)}{src_row}"][0]) is None)
                if s and s not in row_txt:
                    row_txt.append(s)
            warns.append({"유형": "숫자 아닌 값", "파일": fname,
                          "설명": f"숫자 칸에 '{', '.join(row_txt)}'(이)라고 적혀 있어 0으로 처리"})

        # 시군이 합계 칸에 직접 적은 숫자 검산
        for src_row, dst_row, label in ((src_period, form_period, "본기"), (src_cum, form_cum, "누계")):
            calc = {}
            def val_of(col):
                if col in calc:
                    return calc[col]
                return got.get((col, dst_row), 0)
            for tgt, parts, name in CHECKS:
                calc[tgt] = sum(val_of(p) for p in parts)
                raw = cells.get(f"{tgt}{src_row}")
                if raw and not raw[1]:
                    typed = _to_number(raw[0])
                    if typed is not None and typed != calc[tgt]:
                        errors.append({"유형": "합계 불일치", "파일": fname,
                                       "설명": f"[{label}] {name}: 파일에 적힌 합계 {typed:,} ≠ 세부값을 더한 값 {calc[tgt]:,} (차이 {typed - calc[tgt]:,}) — 시군에 확인 후 고친 파일로 다시 올려 주세요"})

        # 본기 > 누계
        over = [col for col in input_cols if col != NOTE_COL
                and got.get((col, form_period), 0) > got.get((col, form_cum), 0)]
        if over:
            warns.append({"유형": "본기 > 누계", "파일": fname,
                          "설명": f"이번 분기 값이 누계보다 큼: {', '.join(over)}열 — 누계 확인 필요"})

    # 합계가 안 맞는 파일이 있으면 어느 쪽이 맞는지 알 수 없으므로 취합본을 만들지 않는다
    if errors:
        return None, log, warns, errors

    # 취합 시트 : 붙은 시군 시트만 더하는 수식
    for r in (form_period, form_cum):
        for col in input_cols:
            if col == NOTE_COL:
                continue
            refs = [f"'{s}'!{col}{r}" for s in sheet_names]
            total[f"{col}{r}"] = "=" + "+".join(refs) if refs else 0
    total["A1"] = f"토지거래계약 허가 현황(경북 {year}년  {quarter}분기)"

    del wb["서식"]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), log, warns, errors


# ── 화면 ──────────────────────────────────────────────────────────────
def render():
    st.caption("시군에서 받은 토지거래계약 허가 현황 파일을 올리면, 빈 서식에서 취합본을 새로 만듭니다.")
    st.info("📌 결과는 맨 앞 '취합' 시트 + 올린 시군만 직제순(포항시 남구 → … → 울릉군)으로 붙습니다. "
            "파일명이나 제목에 시군명이 있으면 자동 인식합니다(포항은 남구·북구까지).")

    files = st.file_uploader("시군 파일 업로드 (여러 개)", type=["xlsx"],
                             accept_multiple_files=True, key="t18_files")

    # 파일 제목에서 연도·분기를 찾아 기본값으로
    det = {}
    for f in files or []:
        try:
            q = _quarter(read_first_sheet(f.getvalue()).get("A1", ("", False))[0])
            if q:
                det[q] = det.get(q, 0) + 1
        except Exception:
            pass
    dy, dq = max(det, key=det.get) if det else (2026, 1)
    # 위젯은 처음 값만 기억하므로, 올린 파일에서 찾은 연도·분기가 바뀌면 직접 넣어 준다
    if det and st.session_state.get("t18_det") != (dy, dq):
        st.session_state["t18_det"] = (dy, dq)
        st.session_state["t18_year"] = dy
        st.session_state["t18_q"] = dq
    c1, c2 = st.columns(2)
    year = c1.number_input("연도", 2020, 2100, dy, key="t18_year")
    quarter = c2.selectbox("분기", [1, 2, 3, 4], index=dq - 1, key="t18_q")

    if files and st.button("🚀 취합본 만들기", key="t18_go"):
        st.session_state["t18_result"] = (
            *build([(f.name, f.getvalue()) for f in files], int(year), int(quarter)),
            int(year), int(quarter))

    res = st.session_state.get("t18_result")
    if not res:
        return
    out, log, warns, errors, y, q = res

    if errors:
        # 합계 불일치 → 빨간 오류, 취합본 없음
        st.error(f"❌ 합계가 맞지 않는 항목 {len(errors)}건 — 취합본을 만들지 않았습니다. "
                 "아래 파일을 시군에 확인해 고친 뒤 다시 올려 주세요.")
        st.dataframe(errors, use_container_width=True)
        if warns:
            st.warning(f"⚠️ 그 밖에 확인이 필요한 항목 {len(warns)}건")
            st.dataframe(warns, use_container_width=True)
        return

    st.success(f"취합본을 만들었습니다. 시군 {len(log)}곳 · 맨 앞 '취합' 시트 포함 {len(log) + 1}장")
    # 다운로드하면 결과를 지워서(휘발성) 화면에 남지 않게 한다
    st.download_button("📥 취합본 다운로드", out,
                       f"{y}년 {q}분기 토지거래허가 분기보고서(취합).xlsx", key="t18_dl",
                       on_click=lambda: st.session_state.pop("t18_result", None))
    if warns:
        st.warning(f"⚠️ 확인이 필요한 항목 {len(warns)}건 (결과는 정상 생성됨)")
        st.dataframe(warns, use_container_width=True)
    else:
        st.info("특이사항 없이 정상적으로 취합되었습니다.")
    with st.expander("처리 내역 (시트 순서)"):
        st.dataframe(log, use_container_width=True)
